"""
SignatureLab — Academic Report Generator
Generates PDF and DOCX reports summarizing the project methodology,
linear algebra formulations, solved numerical examples, and comparison results.
"""

import io
import os
from typing import Dict, Any, Optional
from datetime import datetime

# PDF Generation via ReportLab
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    HRFlowable,
    KeepTogether,
)

# DOCX Generation via python-docx
try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


def generate_pdf_report(
    comparison_data: Optional[Dict[str, Any]] = None,
    student_details: Optional[Dict[str, str]] = None
) -> bytes:
    """
    Generates a PDF academic project report.
    Returns bytes suitable for downloading via Streamlit or saving to disk.
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45,
    )
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0f172a'),
        alignment=1, # Center
        spaceAfter=6
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#2563eb'),
        alignment=1,
        spaceAfter=15
    )
    
    h1_style = ParagraphStyle(
        'H1',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=17,
        textColor=colors.HexColor('#1e293b'),
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True
    )
    
    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor('#334155'),
        spaceAfter=6
    )
    
    callout_style = ParagraphStyle(
        'Callout',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor('#475569'),
    )
    
    meta = student_details or {
        "institution": "[University Name / College of Engineering]",
        "department": "[Department of Computer Science / Mathematics]",
        "course": "Linear Algebra (Undergraduate)",
        "faculty": "[Faculty Advisor / Course Instructor]",
        "students": "Student 1 (ID: XXXXX), Student 2 (ID: YYYYY), Student 3 (ID: ZZZZZ)",
        "date": datetime.now().strftime("%B %d, %Y")
    }

    elements = []
    
    # Header Title
    elements.append(Paragraph("HANDWRITTEN SIGNATURE VERIFICATION USING LINEAR ALGEBRA", title_style))
    elements.append(Paragraph("An Application of Vector Spaces, Matrices, Inner Products, Norms, and Linear Maps", subtitle_style))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#2563eb'), spaceAfter=12))
    
    # Metadata Table
    meta_data = [
        [Paragraph("<b>Course:</b>", body_style), Paragraph(meta.get("course", "Linear Algebra"), body_style)],
        [Paragraph("<b>Institution:</b>", body_style), Paragraph(meta.get("institution", "University Name"), body_style)],
        [Paragraph("<b>Department:</b>", body_style), Paragraph(meta.get("department", "Mathematics / CS"), body_style)],
        [Paragraph("<b>Supervisor / Faculty:</b>", body_style), Paragraph(meta.get("faculty", "Course Professor"), body_style)],
        [Paragraph("<b>Project Team:</b>", body_style), Paragraph(meta.get("students", "Students"), body_style)],
        [Paragraph("<b>Date of Submission:</b>", body_style), Paragraph(meta.get("date", "October 2026"), body_style)],
    ]
    meta_table = Table(meta_data, colWidths=[130, 390])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    elements.append(meta_table)
    elements.append(Spacer(1, 10))
    
    # Section 1: Abstract & Objectives
    elements.append(Paragraph("1. Abstract & Academic Purpose", h1_style))
    abstract_text = (
        "This laboratory report demonstrates how fundamental linear algebra concepts—namely finite-dimensional "
        "vector spaces, matrix flattening, the Euclidean inner product, vector norms, Cauchy-Schwarz geometry, and "
        "cosine similarity—can be systematically applied to compare handwritten signatures. The study demonstrates "
        "both the mathematical elegance and the physical limitations of direct inner-product similarity when applied to "
        "unconstrained biological biometric artifacts."
    )
    elements.append(Paragraph(abstract_text, body_style))
    
    # Section 2: Mathematical Framework
    elements.append(Paragraph("2. Mathematical Framework", h1_style))
    math_text = (
        "<b>Vector Space Isomorphism:</b> An image with dimensions m × n pixels resides in matrix space "
        "M_{m,n}(R). By standard lexicographic vectorization, vec: M_{m,n}(R) -> R^(mn), the image is mapped "
        "isomorphically to an mn-dimensional real inner product space.<br/><br/>"
        "<b>Inner Product (Dot Product):</b> For vectors u, v in R^(mn):<br/>"
        "&nbsp;&nbsp;&nbsp;&nbsp;&lt;u, v&gt; = u · v = &Sigma; (u_i &times; v_i)<br/><br/>"
        "<b>Euclidean Norm (L2 Norm):</b><br/>"
        "&nbsp;&nbsp;&nbsp;&nbsp;||u||_2 = &radic;&lt;u, u&gt; = &radic;(&Sigma; u_i&sup2;)<br/><br/>"
        "<b>Euclidean Metric (Distance):</b><br/>"
        "&nbsp;&nbsp;&nbsp;&nbsp;d(u, v) = ||u - v||_2 = &radic;(&Sigma; (u_i - v_i)&sup2;)<br/><br/>"
        "<b>Cosine Similarity (Directional Alignment):</b><br/>"
        "&nbsp;&nbsp;&nbsp;&nbsp;cos(&theta;) = (u · v) / (||u||_2 &times; ||v||_2)<br/>"
        "By the Cauchy-Schwarz Inequality, |u · v| &le; ||u||_2 ||v||_2, ensuring -1 &le; cos(&theta;) &le; 1."
    )
    elements.append(Paragraph(math_text, body_style))
    
    # Section 3: Solved Numerical Example
    elements.append(Paragraph("3. Solved Numerical Verification Example", h1_style))
    elements.append(Paragraph(
        "To verify analytical engine calculations independently, consider benchmark vectors A = (1, 2, 3) and B = (2, 3, 4) in R&sup3;:",
        body_style
    ))
    
    calc_rows = [
        [Paragraph("<b>Operation</b>", body_style), Paragraph("<b>Formula & Analytical Step</b>", body_style), Paragraph("<b>Exact Result</b>", body_style)],
        [Paragraph("Dot Product", body_style), Paragraph("A · B = (1)(2) + (2)(3) + (3)(4) = 2 + 6 + 12", body_style), Paragraph("20", body_style)],
        [Paragraph("Norm of A", body_style), Paragraph("||A||_2 = &radic;(1&sup2; + 2&sup2; + 3&sup2;) = &radic;(1 + 4 + 9)", body_style), Paragraph("&radic;14 &asymp; 3.7417", body_style)],
        [Paragraph("Norm of B", body_style), Paragraph("||B||_2 = &radic;(2&sup2; + 3&sup2; + 4&sup2;) = &radic;(4 + 9 + 16)", body_style), Paragraph("&radic;29 &asymp; 5.3852", body_style)],
        [Paragraph("Euclidean Distance", body_style), Paragraph("d(A,B) = &radic;((1-2)&sup2; + (2-3)&sup2; + (3-4)&sup2;) = &radic;(1 + 1 + 1)", body_style), Paragraph("&radic;3 &asymp; 1.7321", body_style)],
        [Paragraph("Cosine Similarity", body_style), Paragraph("cos(&theta;) = 20 / (&radic;14 &times; &radic;29) = 20 / &radic;406", body_style), Paragraph("20/&radic;406 &asymp; 0.9926", body_style)],
    ]
    calc_table = Table(calc_rows, colWidths=[120, 270, 130])
    calc_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e2e8f0')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#94a3b8')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(calc_table)
    elements.append(Spacer(1, 10))
    
    # Section 4: Dynamic Comparison Results (if provided)
    if comparison_data:
        elements.append(Paragraph("4. Experimental Signature Comparison Results", h1_style))
        exp_rows = [
            [Paragraph("<b>Metric</b>", body_style), Paragraph("<b>Computed Numerical Value</b>", body_style)],
            [Paragraph("Canvas Dimensions (m &times; n)", body_style), Paragraph(str(comparison_data.get("canvas_shape", "128 x 256")), body_style)],
            [Paragraph("Vector Space Dimension (R^N)", body_style), Paragraph(f"{comparison_data.get('dimension', 32768):,} components", body_style)],
            [Paragraph("Dot Product (u · v)", body_style), Paragraph(f"{comparison_data.get('dot_product', 0.0):.4f}", body_style)],
            [Paragraph("Norm of Reference Vector ||u||_2", body_style), Paragraph(f"{comparison_data.get('norm_u', 0.0):.4f}", body_style)],
            [Paragraph("Norm of Test Vector ||v||_2", body_style), Paragraph(f"{comparison_data.get('norm_v', 0.0):.4f}", body_style)],
            [Paragraph("Euclidean Distance d(u, v)", body_style), Paragraph(f"{comparison_data.get('euclidean_distance', 0.0):.4f}", body_style)],
            [Paragraph("Cosine Similarity cos(&theta;)", body_style), Paragraph(f"<b>{comparison_data.get('cosine_similarity', 0.0):.4f}</b>", body_style)],
            [Paragraph("Subtended Subspace Angle (&theta;)", body_style), Paragraph(f"{comparison_data.get('angle_degrees', 0.0):.2f}&deg;", body_style)],
        ]
        exp_table = Table(exp_rows, colWidths=[200, 320])
        exp_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e0e7ff')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#6366f1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#c7d2fe')),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(exp_table)
        elements.append(Spacer(1, 10))

    # Section 5: Academic Disclaimer & Viva Discussion
    elements.append(Paragraph("5. Limitations & Academic Discussion for Viva", h1_style))
    disc_text = (
        "<b>Rigorous Mathematical Distinction:</b><br/>"
        "• Numerical vector similarity (cos &theta;) measures raw geometric overlap in R^(mn).<br/>"
        "• It constitutes a provisional academic matching heuristic, NOT a legal or biometric identity proof.<br/>"
        "• <b>Vulnerabilities:</b> Pure pixel-level vector spaces lack translation and rotation invariance: a genuine signature "
        "shifted horizontally by 5 pixels will exhibit a substantial decrease in cosine similarity despite human visual equivalence.<br/>"
        "• <b>Educational Context:</b> This implementation highlights the boundary between pure linear algebra models and "
        "invariance manifolds in computer vision."
    )
    elements.append(Paragraph(disc_text, body_style))
    elements.append(Spacer(1, 8))
    
    # Mandatory Footer callout
    elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#cbd5e1'), spaceBefore=8, spaceAfter=8))
    footer_text = "<b>Linear Algebra Mini Project — Educational Demonstration</b><br/>" \
                  "Notice: This application and report serve exclusively as an educational study. " \
                  "It does not represent a production-ready authentication system or guarantee forensic authenticity."
    elements.append(Paragraph(footer_text, callout_style))
    
    doc.build(elements)
    pdf_bytes = buf.getvalue()
    buf.close()
    return pdf_bytes


def generate_docx_report(
    comparison_data: Optional[Dict[str, Any]] = None,
    student_details: Optional[Dict[str, str]] = None
) -> bytes:
    """
    Generates a Microsoft Word (.docx) report using python-docx.
    """
    if not DOCX_AVAILABLE:
        raise RuntimeError("python-docx is not installed in the environment.")
        
    doc = docx.Document()
    
    # Title
    title = doc.add_paragraph()
    r_title = title.add_run("HANDWRITTEN SIGNATURE VERIFICATION USING LINEAR ALGEBRA")
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(15, 23, 42)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Subtitle
    sub = doc.add_paragraph()
    r_sub = sub.add_run("An Application of Vector Spaces, Matrices, Inner Products, Norms, and Linear Maps")
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(37, 99, 235)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph("─" * 45).alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Project Identity Info
    meta = student_details or {
        "institution": "[University Name / College of Engineering]",
        "department": "[Department of Computer Science / Mathematics]",
        "course": "Linear Algebra (Undergraduate)",
        "faculty": "[Faculty Advisor / Course Instructor]",
        "students": "Student 1 (ID: XXXXX), Student 2 (ID: YYYYY), Student 3 (ID: ZZZZZ)",
        "date": datetime.now().strftime("%B %d, %Y")
    }
    
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = 'Light Shading Accent 1'
    for k, v in [
        ("Course:", meta.get("course", "Linear Algebra")),
        ("Institution:", meta.get("institution", "")),
        ("Department:", meta.get("department", "")),
        ("Supervisor / Faculty:", meta.get("faculty", "")),
        ("Student Team:", meta.get("students", "")),
        ("Date:", meta.get("date", "")),
    ]:
        row = tbl.add_row()
        row.cells[0].paragraphs[0].add_run(k).bold = True
        row.cells[1].paragraphs[0].add_run(v)
        
    doc.add_heading("1. Abstract & Academic Purpose", level=1)
    doc.add_paragraph(
        "This project demonstrates how linear algebra principles—including high-dimensional vector spaces, "
        "matrix vectorization, Euclidean inner products, norms, and cosine similarity—enable systematic comparison "
        "of handwritten signatures for academic analysis."
    )
    
    doc.add_heading("2. Linear Algebra Formulation", level=1)
    doc.add_paragraph(
        "• Vectorization: M_(m,n)(R) -> R^(mn)\n"
        "• Inner Product: <u, v> = u · v = ∑ (u_i * v_i)\n"
        "• L2 Norm: ||u||_2 = √(∑ u_i²)\n"
        "• Euclidean Metric: d(u, v) = ||u - v||_2 = √(∑ (u_i - v_i)²)\n"
        "• Cosine Similarity: cos(θ) = (u · v) / (||u||_2 * ||v||_2)\n"
        "• Cauchy-Schwarz Theorem: |u · v| ≤ ||u||_2 ||v||_2, with equality iff u and v are linearly dependent."
    )
    
    doc.add_heading("3. Solved Numerical Verification Example", level=1)
    doc.add_paragraph(
        "Given benchmark vectors A = (1, 2, 3) and B = (2, 3, 4) in R³:\n"
        "• Dot Product: A · B = 1(2) + 2(3) + 3(4) = 20\n"
        "• Norm ||A||: √(1 + 4 + 9) = √14 ≈ 3.7417\n"
        "• Norm ||B||: √(4 + 9 + 16) = √29 ≈ 5.3852\n"
        "• Euclidean Distance d(A,B): √(1 + 1 + 1) = √3 ≈ 1.7321\n"
        "• Cosine Similarity cos(θ): 20 / (√14 * √29) = 20 / √406 ≈ 0.9926"
    )
    
    if comparison_data:
        doc.add_heading("4. Experimental Results", level=1)
        res_tbl = doc.add_table(rows=0, cols=2)
        res_tbl.style = 'Light Grid Accent 1'
        for k, v in [
            ("Canvas Dimensions", str(comparison_data.get("canvas_shape", "128 x 256"))),
            ("Vector Space Dimension", f"{comparison_data.get('dimension', 32768):,}"),
            ("Dot Product", f"{comparison_data.get('dot_product', 0.0):.4f}"),
            ("Norm ||u|| (Ref)", f"{comparison_data.get('norm_u', 0.0):.4f}"),
            ("Norm ||v|| (Test)", f"{comparison_data.get('norm_v', 0.0):.4f}"),
            ("Euclidean Distance", f"{comparison_data.get('euclidean_distance', 0.0):.4f}"),
            ("Cosine Similarity", f"{comparison_data.get('cosine_similarity', 0.0):.4f}"),
            ("Subtended Angle θ", f"{comparison_data.get('angle_degrees', 0.0):.2f}°"),
        ]:
            row = res_tbl.add_row()
            row.cells[0].paragraphs[0].add_run(k).bold = True
            row.cells[1].paragraphs[0].add_run(v)
            
    doc.add_heading("5. Limitations & Viva Preparation", level=1)
    doc.add_paragraph(
        "1. Linear algebra metrics measure exact spatial overlap. They do not possess inherent invariance to translations, "
        "rotations, or non-rigid deformations.\n"
        "2. High cosine similarity indicates close directional orientation in R^(mn), but does NOT guarantee biometric authenticity.\n"
        "3. Thresholding (Otsu binarization) is a non-linear operation, whereas dot products and norms are multilinear / norm functions."
    )
    
    doc.add_paragraph()
    footer_p = doc.add_paragraph()
    r_foot = footer_p.add_run(
        "Linear Algebra Mini Project — Educational Demonstration\n"
        "This project is strictly an educational demonstration and does not guarantee authenticity."
    )
    r_foot.font.size = Pt(9)
    r_foot.font.italic = True
    r_foot.font.color.rgb = RGBColor(100, 116, 139)
    
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()
    buf.close()
    return docx_bytes
